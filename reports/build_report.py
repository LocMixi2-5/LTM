from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH as Align
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
import subprocess
import sys
base=Path(__file__).resolve().parent
full='--full' in sys.argv
stem='Bao_cao_tong_hop_nhom_LTM' if full else 'Bao_cao_LTM_4_phan'
title='Báo cáo tổng hợp đồ án — Trò chơi bi-a hai người qua mạng' if full else 'Báo cáo LTM — 31 use case và 14 biểu đồ UML'
figure_count=21 if full else 14
front_titles=['TÓM TẮT','PHÂN CÔNG MỚI VÀ QUY ƯỚC','PHÂN CÔNG VÀ QUY ƯỚC','DANH MỤC 14 BIỂU ĐỒ','DANH MỤC ẢNH MINH HỌA GIAO DIỆN','TÀI LIỆU VÀ NGUỒN ĐỐI CHIẾU']
if full:front_titles=['TÓM TẮT','DANH MỤC HÌNH']
p=base/(stem+'.docx')
subprocess.run([sys.executable,str(base/('write_group_report.py' if full else 'write_uml_report.py'))],check=True)
subprocess.run(['pandoc',str(base/(stem+'.md')),'--resource-path',str(base),'-o',str(p)],check=True)
d=Document(p)
styles={s.name:s for s in d.styles}
for s in d.sections:
 s.page_width=Cm(21); s.page_height=Cm(29.7)
 s.top_margin=Cm(2.2);s.bottom_margin=Cm(2.2);s.left_margin=Cm(2.7);s.right_margin=Cm(2)
 s.header_distance=Cm(1);s.footer_distance=Cm(1)
 s.different_first_page_header_footer=True
 h=s.header.paragraphs[0];h.alignment=Align.RIGHT
 r=h.add_run('LẬP TRÌNH MẠNG • 8-BALL POOL');r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('666666')
 f=s.footer.paragraphs[0];f.alignment=Align.CENTER
 f.add_run('Trang ')
 field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
 for r in f.runs:r.font.size=Pt(10)
for name in ['Normal','Body Text','First Paragraph']:
 st=styles[name];st.font.name='Times New Roman';st.font.size=Pt(12)
 st.paragraph_format.line_spacing=1.25;st.paragraph_format.space_after=Pt(6)
 st.paragraph_format.widow_control=True
 st.paragraph_format.alignment=Align.JUSTIFY
 st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Times New Roman')
for name,size,color in [('Heading 1',16,'17365D'),('Heading 2',13,'234E70'),('Heading 3',12,'234E70')]:
 st=styles[name];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.bold=True;st.font.color.rgb=RGBColor.from_string(color)
 st.paragraph_format.keep_with_next=True;st.paragraph_format.space_before=Pt(12);st.paragraph_format.space_after=Pt(7)
 st.paragraph_format.alignment=Align.LEFT
for par in d.paragraphs:
 text=par.text
 if text.startswith('PHẦN ') or (full and par.style.name=='Heading 1'):par.paragraph_format.page_break_before=True
 if text in front_titles:
  par.paragraph_format.page_break_before=True
 if text.startswith('Hình '):
  par.alignment=Align.CENTER
  par.paragraph_format.space_before=Pt(6);par.paragraph_format.space_after=Pt(8)
  for r in par.runs:r.font.italic=True;r.font.size=Pt(10.5);r.font.color.rgb=RGBColor.from_string('666666')
 if text.startswith('Thành viên phụ trách:'):
  for r in par.runs:r.font.bold=True
for i,par in enumerate(d.paragraphs):
 if par.text=='TÓM TẮT':break
 par.alignment=Align.CENTER
 if i==0:
  par.paragraph_format.space_before=Pt(48);par.paragraph_format.space_after=Pt(24)
  for r in par.runs:r.font.size=Pt(19);r.font.bold=True
 elif i==1:
  par.paragraph_format.space_after=Pt(30)
  for r in par.runs:r.font.size=Pt(17);r.font.bold=True
 else:par.paragraph_format.space_after=Pt(12)
for table in d.tables:
 table.style=styles['Table'];table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
 n=len(table.columns)
 widths={3:[3.9,5.8,5.6],4:[1.1,3.3,4.7,6.2],5:[1.3,3.9,3.4,1.2,5.5]}.get(n,[15.3/n]*n)
 if n==3 and table.cell(0,0).text=='Mã':widths=[1.2,9.2,4.9]
 for col,w in zip(table.columns,widths):col.width=Cm(w)
 borders=OxmlElement('w:tblBorders')
 for side in ['top','left','bottom','right','insideH','insideV']:
  e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'B8C4CE');borders.append(e)
 table._tbl.tblPr.append(borders)
 for ri,row in enumerate(table.rows):
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  if ri==0:pr.append(OxmlElement('w:tblHeader'))
  for ci,cell in enumerate(row.cells):
   cell.width=Cm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   if ri==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E6EEF5');cell._tc.get_or_add_tcPr().append(sh)
   for par in cell.paragraphs:
    par.alignment=Align.LEFT;par.paragraph_format.line_spacing=1.05
    par.paragraph_format.space_before=Pt(3);par.paragraph_format.space_after=Pt(3)
    par.paragraph_format.keep_with_next=(ri==0)
    for r in par.runs:r.font.name='Times New Roman';r.font.size=Pt(10.5);r.font.bold=(ri==0)
anchor=next(p for p in d.paragraphs if p.text=='TÓM TẮT') if full else next(p for p in d.paragraphs if p.text.startswith('PHẦN 1.'))
head=anchor.insert_paragraph_before('MỤC LỤC NỘI DUNG' if full else 'NỘI DUNG CHÍNH',styles['Heading 2'])
head.paragraph_format.page_break_before=True
contents=[x for x in d.paragraphs if (x.style.name=='Heading 1' and x.text!='BÁO CÁO TỔNG HỢP ĐỒ ÁN') or x.text in front_titles] if full else []
bookmark_id=max([int(x.get(qn('w:id'))) for x in d._element.xpath('.//w:bookmarkStart')]+[0])+1
for section_index,section in enumerate(contents,1):
 line=anchor.insert_paragraph_before()
 line.paragraph_format.space_after=Pt(5)
 line.paragraph_format.line_spacing=1.0
 line.alignment=Align.LEFT
 mark_name='group_section_'+str(section_index)
 mark=OxmlElement('w:bookmarkStart');mark.set(qn('w:id'),str(bookmark_id));mark.set(qn('w:name'),mark_name);section._p.insert(1,mark)
 end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(bookmark_id));section._p.append(end);bookmark_id+=1
 link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'),mark_name)
 run=OxmlElement('w:r');props=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'225F8B');props.append(color);run.append(props)
 text=OxmlElement('w:t');text.text=section.text;run.append(text);link.append(run);line._p.append(link)
for section_title in ([] if full else [
 'Phần 1. TV1 — Quản lý kết nối & Giao thức truyền tin',
 'Phần 2. TV2 — Xác thực & Quản lý dữ liệu',
 'Phần 3. TV3 — Gameplay & Engine Vật lý',
 'Phần 4. TV4 — Giao diện & Cá nhân hóa']):
 anchor.insert_paragraph_before(section_title)
d.core_properties.title=title
d.core_properties.subject='Báo cáo hệ thống thống nhất ngày 29/10/2026, đối chiếu mã 00bf1cc' if full else 'Phân công TV1–TV4 ngày 28/10/2026, đối chiếu mã 00bf1cc'
d.core_properties.author='Nhóm đồ án LTM'
d.core_properties.keywords='LTM, JavaFX, TCP, PostgreSQL, trò chơi qua mạng'
d.core_properties.comments='Nội dung theo quy trình sử dụng; đặc tả và UML ở phụ lục. Mã ứng dụng 00bf1cc.' if full else 'Phân công mới: 7/7/8/9 use case và 3/4/3/4 biểu đồ. Mã ứng dụng 00bf1cc.'
# Keep figures inside the printable area and captions with their picture.
for shape in d.inline_shapes:
 scale=min(1.0,int(Cm(15.3))/shape.width,int(Cm(20.5))/shape.height)
 shape.width=int(shape.width*scale);shape.height=int(shape.height*scale)
for par in d.paragraphs:
 if par.text.startswith('Phạm vi:'):par.paragraph_format.keep_with_next=True
 if par._p.xpath('.//w:drawing'):
  par.alignment=Align.CENTER
  par.paragraph_format.keep_with_next=True
  par.paragraph_format.space_before=Pt(8)
  par.paragraph_format.space_after=Pt(3)
 if par.style.name in ['Source Code','Block Text']:
  par.paragraph_format.alignment=Align.LEFT
  for r in par.runs:r.font.name='DejaVu Sans Mono';r.font.size=Pt(8)
if full:
 for previous,current in zip(d.paragraphs,d.paragraphs[1:]):
  if current._p.xpath('.//w:drawing'):previous.paragraph_format.keep_with_next=True
d.save(p)
r=Document(p)
parts=[x.text for x in r.paragraphs if x.style.name=='Heading 1' and x.text.startswith('CHƯƠNG ' if full else 'PHẦN ')]
assert len(parts)==(7 if full else 4),parts
assert len(r.inline_shapes)==figure_count
assert len(r.tables)>=7,len(r.tables)
body='\n'.join(x.text for x in r.paragraphs)
for name in ['TV1','TV2','TV3','TV4']:assert (name not in body) if full else (name in body)
assert 'không cần tạo ảnh' not in body
print('PASS:',len(parts),'main sections; 14 UML diagrams;',figure_count-14,'UI figures.')
print('Paragraphs:',len(r.paragraphs),'Figures:',len(r.inline_shapes))
print('DOCX size:',p.stat().st_size)
print('Markdown words:',len((base/(stem+'.md')).read_text().split()))

# Optional PDF and standalone HTML, using the same Markdown and image files.
if '--pdf' in sys.argv:
 from lxml import html,etree
 html_path=base/(stem+'.html')
 subprocess.run(['pandoc',str(base/(stem+'.md')),'--standalone','--metadata','title='+title,'--css=report.css','--output',str(html_path)],check=True)
 doc=html.document_fromstring(html_path.read_text())
 body=doc.find('body');children=list(body)
 start=next(e for e in children if e.tag=='h1')
 cover=html.Element('div',attrib={'class':'cover'})
 body.insert(body.index(start),cover)
 for node in children[children.index(start):]:
  if node.tag=='h2' and node.text_content().strip()=='TÓM TẮT':break
  cover.append(node)
 for h in doc.xpath('//h2'):
  if ' '.join(h.text_content().split()) in front_titles:h.set('class','front-break')
 if full:
  toc=html.Element('nav',attrib={'class':'report-toc'})
  heading=etree.SubElement(toc,'h2');heading.text='MỤC LỤC'
  for h in doc.xpath('//body/h1 | //body/h2'):
   if h.tag=='h1' or ' '.join(h.text_content().split()) in front_titles:
    h.set('id','report-section-'+str(len(toc)))
    item=etree.SubElement(toc,'p');link=etree.SubElement(item,'a',href='#'+h.get('id'));link.text=' '.join(h.text_content().split())
  body.insert(body.index(cover)+1,toc)
 for paragraph in doc.xpath('//p'):
  if paragraph.text_content().startswith('Phạm vi:'):paragraph.set('class','scope')
  if full and paragraph.getnext() is not None and paragraph.getnext().tag=='figure':paragraph.set('class','scope')
 html_path.write_text(etree.tostring(doc,encoding='unicode',method='html',doctype='<!DOCTYPE html>'))
 subprocess.run([sys.executable,'-m','weasyprint',str(html_path),str(base/(stem+'.pdf'))],check=True)
