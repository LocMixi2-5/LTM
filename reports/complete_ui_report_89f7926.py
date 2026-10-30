"""Fill the deployment UI section of the revised upload with current captures."""
from pathlib import Path
from shutil import copy2
import json
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image
base=Path(__file__).resolve().parent
images=base/'images_89f7926'
for p in Path('/workspace/ltm-env/report-89f7926/ui-images').glob('*.png'):copy2(p,images/p.name)
items=[
('Đăng nhập','login.png','Màn hình mở đầu nhận tên đăng nhập hoặc email và mật khẩu. Người dùng có thể chuyển sang đăng ký hoặc mở cấu hình kết nối trước khi vào sảnh.','LoginScene'),
('Cấu hình kết nối server','server_config.png','Hộp thoại cho phép nhập địa chỉ server và cổng kết nối. Đây là bước cấu hình đầu vào khi hai máy client cần kết nối đến cùng một server.','ServerConfigDialog'),
('Đăng ký tài khoản','registration_actual.png','Form gồm email, tên đăng nhập, mật khẩu và xác nhận mật khẩu. Dữ liệu hợp lệ mới được gửi đi; server tiếp tục kiểm tra cú pháp và thông tin trùng lặp.','RegisterScene'),
('Phản hồi dữ liệu đăng ký không hợp lệ','registration_validation.png','Thông báo lỗi xuất hiện ngay trên form khi email sai cú pháp. Trạng thái này giúp người dùng sửa đầu vào mà không cần rời màn hình đăng ký.','RegisterScene'),
('Sảnh chính','lobby.png','Sảnh tập trung thông tin tài khoản, điểm kinh nghiệm, cấp độ và danh sách bạn. Từ đây người dùng truy cập tìm trận, luyện tập, bảng xếp hạng và cài đặt.','LobbyScene'),
('Menu thao tác với bạn bè','friend_menu_actual.png','Menu cạnh tên bạn cung cấp các thao tác tương tác. Giao diện giữ menu đang mở khi danh sách tự cập nhật và áp dụng dữ liệu mới sau khi đóng.','LobbyScene'),
('Tìm người chơi và lời mời kết bạn','friend_search.png','Phần trên hiển thị kết quả tìm kiếm và trạng thái đã là bạn hoặc nút kết bạn. Phần dưới tổng hợp lời mời đang chờ với nút đồng ý và từ chối.','FriendSearchDialog'),
('Hội thoại với bạn bè','chat_actual.png','Khung hội thoại hiển thị người nhận, lịch sử, nội dung tin và thời gian. Bản hiện tại kiểm tra tên người nhận trước khi áp dụng lịch sử và giữ các tin cố ý giống nội dung.','LobbyScene'),
('Nhận lời mời thách đấu','challenge_actual.png','Popup hiển thị người mời, lựa chọn phản hồi và thời gian chờ. Bộ đếm giao diện là 10 giây; thời hạn lời mời phía server là 15 giây.','LobbyScene'),
('Thông báo kết quả lời mời','challenge_toast.png','Thông báo ngắn trên sảnh phản hồi kết quả lời mời. Toast không mở thêm hộp thoại buộc người dùng đóng trước khi tiếp tục thao tác.','LobbyScene'),
('Tìm trận','matchmaking.png','Màn hình thể hiện trạng thái chờ ghép đối thủ và cung cấp nút hủy. Khi nhận thông tin ghép trận, client chuyển sang màn hình chuẩn bị.','MatchmakingScene'),
('Chuẩn bị vào trận','pregame.png','Hai người chơi được trình bày đối diện nhau cùng thông tin cấp độ và trạng thái chuẩn bị. Client gửi tín hiệu sẵn sàng; server yêu cầu hai người hợp lệ cùng sẵn sàng để bắt đầu.','PreGameScene'),
('Bàn chơi với đối thủ','game_actual.png','Màn hình gồm bàn bi, thông tin hai người chơi, lượt, thời gian, nhóm bi và bộ điều khiển xoáy. Hình minh họa client nhận mẫu ngắm của đối thủ; không dùng ảnh đơn để kết luận độ mượt đồng bộ hai máy.','GameScene'),
('Bàn luyện tập','practice_actual.png','Chế độ luyện tập dùng bàn chơi và cơ chế điều khiển tương tự thi đấu nhưng ẩn phần thông tin đối thủ. Người chơi có thể luyện đặt bi, ngắm và đánh bi trước khi tham gia trận.','GameScene'),
('Cài đặt avatar','settings_avatar.png','Tab Avatar cho phép xem trước và chọn hình đại diện. Các nút đóng và lưu nằm ở cuối hộp thoại; thay đổi hình đại diện được dùng khi hiển thị người chơi.','SettingsDialog'),
('Bộ sưu tập gậy','settings_cues.png','Tab bộ sưu tập hiển thị các gậy, vùng xem trước và điều kiện sử dụng theo cấp độ. Người chơi xem trước hoặc trang bị gậy đủ điều kiện; server kiểm tra lại cấp khi nhận cập nhật.','SettingsDialog'),
('Âm thanh và đồ họa','settings_audio.png','Tab tập hợp các điều khiển âm thanh và hiệu ứng đang có trong ứng dụng. Người dùng điều chỉnh trải nghiệm nghe và hiển thị trong cùng hộp thoại cài đặt.','SettingsDialog'),
('Bảng xếp hạng','leaderboard.png','Bảng hiển thị thứ hạng, tên, cấp độ, kinh nghiệm, số trận thắng và thua; hàng của người đang đăng nhập được phân biệt. Dữ liệu trong ảnh là dữ liệu minh họa có kiểm soát.','LeaderboardScene'),
('Kết quả chiến thắng','result_win.png','Màn hình thông báo chiến thắng, cặp đấu và kinh nghiệm nhận được, kèm nút trở về sảnh. Bản 89f7926 không còn khung thông báo mở khóa skin trong màn hình này.','ResultScene'),
('Kết quả thất bại','result_loss.png','Nhánh thất bại dùng cùng bố cục kết quả, thay đổi nội dung và điểm nhận được. Hai ảnh kết quả được mở trực tiếp bằng dữ liệu minh họa, không phải kết quả của một ván đầu cuối đã nghiệm thu.','ResultScene'),
]
file=base/'Bao_Cao_LTM_cap_nhat_89f7926.docx'
doc=Document(file)
start=next(p for p in doc.paragraphs if p.text=='2. Giao diện')._p
end=next(p for p in doc.paragraphs if p.text=='3. Các sự cố triển khai thường gặp')._p
node=start.getnext()
while node is not end:
 nxt=node.getnext();node.getparent().remove(node);node=nxt

def add(text='',style='Body Text'):
 p=doc.add_paragraph(text,style);end.addprevious(p._p);return p
add('Các hình sau bao quát toàn bộ các lớp màn hình và hộp thoại của client: đăng nhập, đăng ký, sảnh, tìm trận, chuẩn bị, chơi/luyện tập, kết quả, bảng xếp hạng và ba hộp thoại cấu hình/kết bạn/cài đặt. Các trạng thái popup, toast, lỗi nhập liệu và ba tab cài đặt được minh họa riêng.')
add('Nguồn: chạy trực tiếp các lớp JavaFX từ mã 89f7926 trên Linux/Xvfb ngày 30/10/2026; tên tài khoản và dữ liệu mạng được giả lập để chụp. Ảnh phản ánh giao diện của bản đối chiếu, không thay cho nghiệm thu EXE Windows hoặc một ván hoàn chỉnh giữa hai máy. Chỉ giữ tên chức năng có sẵn của ứng dụng; không thêm nhãn thành viên lên ảnh.')
for i,(title,name,desc,cls) in enumerate(items,1):
 h=add(f'2.{i}. {title}','Heading 3');h.paragraph_format.keep_with_next=True
 p=add(desc);p.paragraph_format.keep_with_next=True
 p=add();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(4)
 w,ht=Image.open(images/name).size
 width=min(16.0,13.2*w/ht)
 p.add_run().add_picture(str(images/name),width=Cm(width))
 cap=add(f'Hình 8.{i} — {title}.','Caption');cap.alignment=WD_ALIGN_PARAGRAPH.CENTER;cap.paragraph_format.keep_with_next=False
# Update provenance where the previous pass retained historical screenshots.
for p in doc.paragraphs:
 if p.text.startswith('Minh chứng lịch sử và nguồn ảnh ngày 26/10/2026:'):
  p.text='Nguồn ảnh giao diện hiện hành: reports/images_89f7926/README.md và manifest.json; harness chụp trong reports/evidence/89f7926/. Log 26/10/2026 được giữ ở reports/evidence/00bf1cc/ để đối chiếu lịch sử, không dùng thay kết quả mới.'
 if p.text.startswith('Qua đọc mã, hệ thống còn thiếu token/sequence'):
  p.text=p.text.replace('Các thay đổi phản xạ băng/đường ngắm và ResultScene mới được đối chiếu mã, chưa nghiệm thu trải nghiệm đầy đủ.','Các thay đổi phản xạ băng/đường ngắm mới được đối chiếu mã và biên dịch. ResultScene đã được mở và chụp ở cả hai nhánh thắng/thua với dữ liệu giả lập; chưa nghiệm thu kết quả của ván đầu cuối.')
 if p.text.startswith('ResultScene hiện hiển thị kết quả') and 'Hình 8.19–8.20' not in p.text:
  p.text+=' Hai nhánh thắng/thua được minh họa ở Hình 8.19–8.20.'
doc.save(file)
(images/'manifest.json').write_text(json.dumps([{'figure':f'8.{i}','title':v[0],'file':v[1],'description':v[2],'source_class':v[3]} for i,v in enumerate(items,1)],ensure_ascii=False,indent=2)+'\n')
print('Inserted',len(items),'current screenshots; preserved ten main parts and existing technical figures.')
