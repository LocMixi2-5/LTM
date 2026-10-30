"""Add a linked contents page using page numbers from a rendered PDF."""
from pathlib import Path
import re,fitz
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm,Pt
from docx.enum.text import WD_TAB_ALIGNMENT,WD_TAB_LEADER
b=Path(__file__).resolve().parent
file=b/'Bao_Cao_LTM_cap_nhat_89f7926.docx'
d=Document(file);pdf=fitz.open('/workspace/ltm-env/report-89f7926/Bao_Cao_LTM_cap_nhat_89f7926.pdf')
norm=lambda x:re.sub(r'\s+',' ',x).strip()
headings=[p for p in d.paragraphs if p.style.name=='Heading 1' and p.text!='Mục Lục']
start=next(p for p in d.paragraphs if p.text=='Mục Lục')._p
end=headings[0]._p
for br in start.xpath('.//w:br'):
 if br.get(qn('w:type'))=='page':br.getparent().remove(br)
headings[0].paragraph_format.page_break_before=True
n=start.getnext()
while n is not end:
 nxt=n.getnext();n.getparent().remove(n);n=nxt
for i,h in enumerate(headings):
 pages=[j+1 for j,p in enumerate(pdf) if j!=1 and norm(h.text) in norm(p.get_text())]
 assert pages,(h.text,'missing in PDF')
 page=pages[0];name=f'Report89Section{i}'
 # Re-running retains one bookmark per heading.
 for child in list(h._p):
  if child.tag in [qn('w:bookmarkStart'),qn('w:bookmarkEnd')] and child.get(qn('w:id'))==str(890+i):h._p.remove(child)
 bs=OxmlElement('w:bookmarkStart');bs.set(qn('w:id'),str(890+i));bs.set(qn('w:name'),name)
 be=OxmlElement('w:bookmarkEnd');be.set(qn('w:id'),str(890+i));h._p.insert(1,bs);h._p.append(be)
 p=d.add_paragraph();end.addprevious(p._p)
 p.paragraph_format.space_after=Pt(8);p.paragraph_format.line_spacing=1.1
 p.paragraph_format.tab_stops.add_tab_stop(Cm(16),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
 link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'),name)
 r=OxmlElement('w:r');pr=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),'22');pr.append(sz);r.append(pr)
 t=OxmlElement('w:t');t.text=h.text;r.append(t);link.append(r);p._p.append(link)
 p.add_run('\t'+str(page)).font.size=Pt(11)
 print(page,h.text)
for table in d.tables:
 if table.cell(0,0).text=='Mã' and table.cell(0,1).text=='Kịch bản kiểm tra':
  widths=[1000,2180,4200,1750]
  for c,w in zip(table._tbl.tblGrid.gridCol_lst,widths):c.set(qn('w:w'),str(w))
  for row in table.rows:
   for c,w in zip(row.cells,widths):c._tc.get_or_add_tcPr().find(qn('w:tcW')).set(qn('w:w'),str(w))
for table in d.tables:
 if table.cell(0,0).text=='Mã' and table.cell(0,1).text=='Kịch bản cần chạy':
  widths=[1000,2180,3200,2750]
  for c,w in zip(table._tbl.tblGrid.gridCol_lst,widths):c.set(qn('w:w'),str(w))
  for row in table.rows:
   for c,w in zip(row.cells,widths):c._tc.get_or_add_tcPr().find(qn('w:tcW')).set(qn('w:w'),str(w))
for p in d.paragraphs:
 if p.text.startswith('Ngoại lệ: Không có dữ liệu hoặc lỗi mạng cần thông báo; chưa có nghiệm thu toàn bộ UI bảng xếp hạng'):
  p.text='Ngoại lệ: Bảng trống được hiển thị khi chưa có dữ liệu. Đã mở và chụp bảng xếp hạng với dữ liệu giả lập trong đợt đối chiếu; chưa nghiệm thu luồng cập nhật xếp hạng từ kết quả của một ván đầy đủ.'
 if p.text.startswith('Minh chứng mới nằm tại reports/evidence/89f7926/'):
  p.text=p.text.replace('build.txt, protocol.txt và ui.txt.','build.txt, protocol.txt, ui.txt và ui-capture.txt. Bộ ảnh hiện hành bao quát 20 giao diện/trạng thái; dữ liệu giả lập chỉ dùng để minh họa màn hình.')
for p in d.paragraphs:
 phrase=' Hai nhánh thắng/thua được minh họa ở Hình 8.19–8.20.'
 while p.text.count(phrase)>1:p.text=p.text.replace(phrase+phrase,phrase)
d.save(file)
