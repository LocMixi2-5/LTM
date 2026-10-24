# -*- coding: utf-8 -*-
"""
HỆ THỐNG SINH BÁO CÁO ĐỒ ÁN MÔN LẬP TRÌNH MẠNG
Đề tài: Game Bida 8-Ball Pool Multiplayer qua mạng TCP/IP & UDP
Định dạng xuất: Microsoft Word (.docx) chuẩn học thuật A4
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DIAGRAMS_DIR = os.path.join(BASE_DIR, 'report_assets', 'diagrams')
OUTPUT_FILE = os.path.join(BASE_DIR, 'Bao_cao_Do_an_Lap_trinh_mang_8BallPool.docx')

# Bảng màu chuẩn học thuật (Academic Navy Palette)
C_NAVY_DARK  = RGBColor(27, 54, 93)     # #1B365D - Tiêu đề chính, Chương
C_NAVY_MED   = RGBColor(43, 108, 176)   # #2B6CB0 - Mục cấp 2
C_TEXT_DARK  = RGBColor(45, 55, 72)     # #2D3748 - Thân bài
C_TEXT_MUTED = RGBColor(113, 128, 150)  # #718096 - Chú thích
C_AMBER      = RGBColor(180, 83, 9)     # #B45309 - Placeholder cảnh báo
C_GREEN      = RGBColor(22, 101, 52)    # #166534 - Trạng thái Đạt

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}>'
                      f'<w:top w:w="{top}" w:type="dxa"/>'
                      f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
                      f'<w:left w:w="{left}" w:type="dxa"/>'
                      f'<w:right w:w="{right}" w:type="dxa"/>'
                      f'</w:tcMar>')
    tcPr.append(tcMar)

def add_header_footer(doc):
    for s in doc.sections:
        s.different_first_page_header_footer = True
        
        # Header
        hdr = s.header
        p_hdr = hdr.paragraphs[0]
        p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_hdr = p_hdr.add_run("Báo Cáo Đồ Án Lập Trình Mạng — Game 8-Ball Pool Multiplayer")
        r_hdr.font.name = "Times New Roman"
        r_hdr.font.size = Pt(8.5)
        r_hdr.font.italic = True
        r_hdr.font.color.rgb = C_TEXT_MUTED
        
        # Footer
        ftr = s.footer
        p_ftr = ftr.paragraphs[0]
        p_ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_ftr = p_ftr.add_run("Trang ")
        r_ftr.font.name = "Times New Roman"
        r_ftr.font.size = Pt(9.5)
        r_ftr.font.color.rgb = C_TEXT_MUTED
        
        fld_page = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
        p_ftr._p.append(fld_page)
        
        r_ftr2 = p_ftr.add_run(" / ")
        r_ftr2.font.name = "Times New Roman"
        r_ftr2.font.size = Pt(9.5)
        r_ftr2.font.color.rgb = C_TEXT_MUTED
        
        fld_numpages = parse_xml(r'<w:fldSimple %s w:instr="NUMPAGES"/>' % nsdecls('w'))
        p_ftr._p.append(fld_numpages)

def create_p(doc, text="", space_before=0, space_after=6, line_spacing=1.35, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_indent=Cm(0.8), bold=False, italic=False, color=C_TEXT_DARK, size=13):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if first_indent and align == WD_ALIGN_PARAGRAPH.JUSTIFY:
        p.paragraph_format.first_line_indent = first_indent
    if text:
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = color
    return p

def add_heading1(doc, text):
    doc.add_page_break()
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = C_NAVY_DARK
    return p

def add_heading2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = C_NAVY_MED
    return p

def add_heading3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = C_TEXT_DARK
    return p

def add_figure(doc, img_path, caption_text, width_cm=15.0):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(10)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Cm(width_cm))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(12)
        r_cap = p_cap.add_run(caption_text)
        r_cap.font.name = "Times New Roman"
        r_cap.font.size = Pt(10.5)
        r_cap.font.italic = True
        r_cap.font.bold = True
        r_cap.font.color.rgb = C_TEXT_DARK
    else:
        print(f"Warning: Image file not found: {img_path}")

def add_callout_placeholder(doc, img_code, title, capture_instructions, suggested_file, target_location):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "FEF3C7") # Vàng cam nhạt cảnh báo
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    cell.width = Cm(15.5)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    
    r1 = p.add_run(f"📸 [{img_code}: {title}]\n")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(11)
    r1.font.bold = True
    r1.font.color.rgb = C_AMBER
    
    r2 = p.add_run(f"• Thao tác chụp thực tế: {capture_instructions}\n"
                   f"• Tên file ảnh gợi ý: {suggested_file}\n"
                   f"• Vị trí gắn trong báo cáo: {target_location}")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(10)
    r2.font.italic = True
    r2.font.color.rgb = RGBColor(120, 53, 15)
    
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(6)

def add_code_block(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
    cell.width = Cm(15.5)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    
    r = p.add_run(code_text.strip())
    r.font.name = "Consolas"
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(30, 41, 59)
    
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(6)

def style_table(table, col_widths, headers, data, align_cols=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_tr = table.rows[0]._tr.get_or_add_trPr()
    header_tr.append(parse_xml(r'<w:tblHeader %s/>' % nsdecls('w')))
    
    for row_idx, row in enumerate(table.rows):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(r'<w:cantSplit %s/>' % nsdecls('w')))
        
        for col_idx, cell in enumerate(row.cells):
            cell.width = Cm(col_widths[col_idx])
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            
            align = WD_ALIGN_PARAGRAPH.LEFT
            if align_cols and col_idx < len(align_cols):
                align = align_cols[col_idx]
            p.alignment = align
            
            if row_idx == 0:
                set_cell_background(cell, "1B365D")
                r = p.add_run(headers[col_idx])
                r.font.name = "Times New Roman"
                r.font.size = Pt(9.5)
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
            else:
                bg = "FFFFFF" if row_idx % 2 == 1 else "F8FAFC"
                set_cell_background(cell, bg)
                val = data[row_idx - 1][col_idx]
                r = p.add_run(str(val))
                r.font.name = "Times New Roman"
                r.font.size = Pt(9)
                if str(val) == "Đạt":
                    r.font.bold = True
                    r.font.color.rgb = C_GREEN
                elif str(val) == "Chưa kiểm thử":
                    r.font.italic = True
                    r.font.color.rgb = C_AMBER
                else:
                    r.font.color.rgb = C_TEXT_DARK

def add_table_with_caption(doc, caption_text, col_widths, headers, data, align_cols=None):
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_cap.paragraph_format.space_before = Pt(8)
    p_cap.paragraph_format.space_after = Pt(3)
    p_cap.paragraph_format.keep_with_next = True
    r = p_cap.add_run(caption_text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(10.5)
    r.font.bold = True
    r.font.italic = True
    r.font.color.rgb = C_NAVY_DARK
    
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    style_table(table, col_widths, headers, data, align_cols)
    
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(8)
    return table

def add_native_toc(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run()
    fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> TOC \o "1-3" \h \z \u </w:instrText>' % nsdecls('w'))
    fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

print("make_report.py helper module loaded successfully.")
